"""Score a wave: quote-verification, reader agreement, and agreement with the
human-filled workbook on the fields the workbook happens to carry.

The workbook is the test set and is read only here, at the top layer. Nothing
in this file is ever shown to a reader.
"""
import json, re, sys, os, unicodedata
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build import assemble
from ids import FILE_RX, corrected
from norm import norm
from paths import ROOT, complete_stems
from verify import check

HERE = os.path.dirname(os.path.abspath(__file__))


def prior_summary():
    """The workbook, wherever this clone keeps it.

    release.json records the path the release was built from, which is absolute
    and on one machine. Every clone still has the file, under the release
    directory at the repository root, so look for it there rather than failing
    with a path nobody else can have. Scoring that only runs on the author's
    laptop is scoring that does not run.
    """
    recorded = json.load(open(os.path.join(HERE, "release.json")))["prior_summary"]
    if os.path.exists(recorded):
        return recorded
    name = os.path.basename(recorded)
    for base, dirs, files in os.walk(ROOT):
        dirs.sort()
        if name in files:
            return os.path.join(base, name)
    return recorded


XLSX = prior_summary()

# A workbook date the workbook actually commits to: a year, a month, or a day.
DATEISH = re.compile(r"^\d{4}(-\d{2}(-\d{2})?)?$")


def deaccent(s):
    """Fold accents away: 'Quebec' and 'Quebec' are one answer written twice.

    Only the scorer folds them. `norm` is the canonical fold that verification
    and reader comparison are both built on, and widening it there would
    re-classify conflicts already banked -- a scoring fix is not worth
    reopening the corpus for.
    """
    return "".join(c for c in unicodedata.normalize("NFD", str(s))
                   if unicodedata.category(c) != "Mn")


def untestable(kind, truth):
    """A workbook value that cannot settle anything.

    Some rows carry a project span where the authorization carries an issuance
    date -- "2019-2022 (phased)". Scoring that as a miss blames the extraction
    for a question the workbook never answered, and scoring it as a hit inflates
    the result. It is neither, and is counted on its own.
    """
    return kind == "date" and not DATEISH.match(str(truth).strip()[:10].strip())

# extracted field -> workbook column, and how to compare
MAP = {
    "date_of_issuance":     ("Authorization_Date",   "date"),
    "proponent":            ("Proponent",            "text"),
    "province":             ("Province",             "text"),
    "hadd_destruction_m2":  ("HADD_Destruction_m2",  "num"),
    "hadd_alteration_m2":   ("HADD_Alteration_m2",   "num"),
    "hadd_disturbance_m2":  ("HADD_Disturbance_m2",  "num"),
    "impact_other_units":   ("HADD_Other_Units",     "text"),
}


def truth_rows():
    import openpyxl
    ws = openpyxl.load_workbook(XLSX)["Authorization_Summary"]
    hdr = [c.value for c in ws[1]]
    idx = {h: i + 1 for i, h in enumerate(hdr)}
    out = {}
    for r in range(2, ws.max_row + 1):
        m = FILE_RX.search(str(ws.cell(row=r, column=idx["DFO_File_or_PATH"]).value or ""))
        if m:
            out.setdefault(m.group(0), {h: ws.cell(row=r, column=idx[h]).value for h in hdr})
    return out


def same(kind, truth, got):
    if truth in (None, "") and got in (None, ""):
        return True
    if truth in (None, "") or got in (None, ""):
        return False
    if kind == "num":
        try:
            t, g = float(truth), float(got)
        except (TypeError, ValueError):
            return False
        return abs(t - g) <= max(1.0, t * 0.005)
    if kind == "date":
        # The workbook is routinely coarser than the document: a row carries
        # "2020" or "2020-08" where the page is stamped with a day. That is the
        # same answer recorded to less precision, not a different one, so
        # compare only as far as the workbook commits.
        t = str(truth).strip()[:10].strip()
        return str(got)[:len(t)] == t
    t, g = norm(deaccent(truth)), norm(deaccent(got))
    return t in g or g in t


def main(sandbox, half="tune"):
    """The work order carries the page range and the stem; do not re-derive them.

    `half` names which side of the split to score. It defaults to the tuning
    half: the holdout is scored once, at the end, and a scorer that reaches it
    by default gets run against it by accident long before then.

    Agreement is measured against the row the pipeline actually builds, not
    against one reader. Nine file numbers in this release are several documents
    -- an authorization and its amendments -- and the workbook gives them one
    row each. Scoring each document against that row separately charged the
    amendment with not being the original, and the original with not being the
    amendment; both were counted as misses and neither was one. `assemble`
    merges them the way the workbook does, so that is what gets scored.
    """
    # `assemble` reads the sandbox's own work order, so this reads the same one.
    # An alternative wave was once accepted here and is not any more: it would
    # have verified quotes over one set of documents and measured agreement over
    # another, which is worse than not offering the option.
    orders = json.load(open(f"{sandbox}/wave.json"))
    keep = None
    if half != "all":
        keep = set(json.load(open(os.path.join(HERE, "split.json")))[half])
        # The split names file numbers as the documents carry them, and a row is
        # filed under the corrected one; compare like with like on both sides.
        orders = [o for o in orders if corrected(o["stem"], o["file_number"]) in keep]
    truth = truth_rows()
    ver = Counter()
    agree = Counter()
    skipped = Counter()
    complete = complete_stems(sandbox)
    unread = {corrected(o["stem"], o["file_number"]) for o in orders
              if o["stem"] not in complete}
    detail = []

    # Quote verification stays per reader: it asks whether each reader put its
    # answers where it said it did, which merging would hide.
    for order in orders:
        if order["pass"] != "a" or order["reader"] != "1":
            continue
        stem, sl = order["stem"], order["slice"]
        lo, hi = order["first_page"], order["last_page"]
        for tag in ("a1", "a2"):
            p = f"{sandbox}/out/{stem}.{tag}.json"
            if not os.path.exists(p):
                continue
            for _, verdict, _ in check(json.load(open(p)), sl, lo, hi):
                ver[verdict] += 1

    rows, _, _, _ = assemble(sandbox)
    scored = set()
    for row in rows:
        doc = row["documents"][0]["file_number"]
        if (keep is not None and doc not in keep) or doc not in truth:
            continue
        scored.add(doc)
        for f, (col, kind) in MAP.items():
            t = truth[doc].get(col)
            if t in (None, ""):
                continue                       # workbook has no label -> not a test case
            if untestable(kind, t):
                skipped[f] += 1
                continue
            entries = row["fields"].get(f) or []
            stated = [e for e in entries if e.get("role") == "total"]
            g = (stated or entries or [{}])[0].get("value")
            ok = same(kind, t, g)
            if not ok and kind == "num" and len(entries) > 1:
                # A figure the reader listed in parts still agrees with a labelled
                # total; scoring the first part alone measured the wrong thing.
                parts = [e["value"] for e in entries if isinstance(e.get("value"), (int, float))]
                if parts and (same(kind, t, sum(parts)) or any(same(kind, t, p) for p in parts)):
                    ok, g = True, f"parts{parts}"
            # A cell the pipeline withheld is not a cell it got wrong. The
            # queue is the designed outcome for a disagreement, so counting it
            # as a wrong answer measures the opposite of what the rule intends.
            outcome = "hit" if ok else ("blank" if not entries else "wrong")
            agree[f"{f}:{outcome}"] += 1
            if not ok:
                detail.append((doc, f, str(t)[:34],
                               "(withheld)" if outcome == "blank" else str(g)[:34],
                               "; ".join(d["doc_id"] for d in row["documents"])))

    print("=== quote verification (both readers, all fields) ===")
    tot = sum(ver.values())
    for k, v in ver.most_common():
        print(f"   {k:10s} {v:5d}  {v/tot:5.1%}")

    print(f"\n=== agreement with the human workbook ({len(scored)} file numbers) ===")
    print(f"   {'field':22s} {'agrees':>7s} {'wrong':>6s} {'withheld':>9s}   {'of cells':>8s}")
    tot = Counter()
    for f in MAP:
        h, w, b = (agree.get(f + ":hit", 0), agree.get(f + ":wrong", 0),
                   agree.get(f + ":blank", 0))
        tot.update({"hit": h, "wrong": w, "blank": b})
        if h + w + b:
            print(f"   {f:22s} {h:7d} {w:6d} {b:9d}   {h/(h+w+b):7.0%}")
    n = sum(tot.values())
    if n:
        print(f"   {'TOTAL':22s} {tot['hit']:7d} {tot['wrong']:6d} {tot['blank']:9d}   "
              f"{tot['hit']/n:7.0%}")

    if skipped:
        print("\n=== not testable: the workbook value is not an answer to this field ===")
        for f, n in skipped.most_common():
            print(f"   {f:22s} {n:2d}")
    if unread:
        print(f"\n{len(unread)} file number(s) in this half are not yet read, and are not scored")
    if detail:
        print("\n=== misses ===")
        for d in detail:
            print(f"   {d[0]:16s} {d[1]:22s} workbook={d[2]:36s} extracted={d[3]}")
            if ";" in d[4]:            # only worth saying when the row is a merge
                print(f"   {'':16s} {'':22s} from {d[4]}")


if __name__ == "__main__":
    # score.py <sandbox> [tune|holdout|all]
    args = sys.argv[1:]
    half = args.pop() if args and args[-1] in ("tune", "holdout", "all") else "tune"
    main(args[0], half)
